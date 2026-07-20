import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import sys

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0]
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.bold = True

def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        p.add_run(text.replace(bold_prefix, "", 1))
    else:
        p.add_run(text)
    p.style.font.name = 'Arial'
    p.style.font.size = Pt(11)

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        p.add_run(text.replace(bold_prefix, "", 1))
    else:
        p.add_run(text)
    p.style.font.name = 'Arial'
    p.style.font.size = Pt(11)

def generate_proposal(output_file):
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Project Proposal', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    # 1. Title
    add_heading(doc, '1. Title of the Project')
    add_paragraph(doc, 'Smart Nepal Travel System')

    # 2. Intro
    add_heading(doc, '2. Introduction and Background')
    add_paragraph(doc, 'The Smart Nepal Travel System is an intelligent web-based platform designed to assist tourists in exploring travel destinations and packages across Nepal. Traditionally, tourists face difficulties in finding reliable travel itineraries, getting instant answers to their queries, and making bookings in a centralized manner. Our software improves this process by integrating an AI chatbot for customer support, a personalized recommendation engine, and a seamless booking system, ensuring a hassle-free travel planning experience.')

    # 3. Problem Statement
    add_heading(doc, '3. Problem Statement')
    add_paragraph(doc, 'The existing travel booking process is fragmented. Tourists often have to visit multiple websites to find destinations, compare packages, and seek assistance. Furthermore, there is a lack of intelligent systems that can understand user preferences and provide personalized recommendations or answer queries in real-time. This project aims to solve these inefficiencies by providing a unified, AI-driven travel platform.')

    # 4. Objectives
    add_heading(doc, '4. Objectives')
    add_bullet(doc, 'To computerize and centralize the travel package booking process.')
    add_bullet(doc, 'To implement an AI-powered chatbot for instant, 24/7 customer assistance.')
    add_bullet(doc, 'To provide personalized travel package recommendations using machine learning (Content-Based Filtering).')
    add_bullet(doc, 'To digitize the payment process with an integrated eSewa payment simulation.')
    add_bullet(doc, 'To provide an admin dashboard with advanced data analytics and forecasting.')

    # 5. Scope and Limitations
    add_heading(doc, '5. Scope and Limitations')
    add_bullet(doc, 'Scope: The system includes modules for User Authentication, Destination & Package Browsing, an AI Chatbot (NLP/Gemini), a Recommender System (TF-IDF), Booking Management, and an Admin Dashboard with data analytics (K-Means Clustering and Linear Regression).', 'Scope:')
    add_bullet(doc, 'Limitations: The payment gateway is a simulation for demonstration purposes. Direct integration with live travel agencies and airlines is excluded due to time constraints, and there is no mobile app version at this stage.', 'Limitations:')

    # 6. Feasibility Study
    add_heading(doc, '6. Feasibility Study')
    add_bullet(doc, 'Technical: The project utilizes standard, well-documented web development tools (Python, Flask, scikit-learn, HTML/CSS). The team possesses the required programming skills and hardware to complete it.', 'Technical:')
    add_bullet(doc, 'Operational: The proposed system features an intuitive web interface, making it extremely easy to use for end-users (tourists) and administrators.', 'Operational:')
    add_bullet(doc, 'Economic: The project is highly cost-effective as it relies primarily on open-source technologies (Python, Flask, SQLite/MySQL) and requires no expensive proprietary software.', 'Economic:')

    # 7. System Requirements
    add_heading(doc, '7. System Requirements')
    add_bullet(doc, 'Hardware: Minimum 4GB RAM, Intel Core i3 or equivalent processor, and 1GB of free storage.', 'Hardware:')
    add_bullet(doc, 'Software: Windows 10/11, macOS, or Linux operating system. Python 3.8+, a modern web browser, VS Code IDE, and a relational database (SQLite/MySQL).', 'Software:')

    # 8. Methodology
    add_heading(doc, '8. Proposed Methodology')
    add_bullet(doc, 'Requirement Collection: Data will be gathered by reviewing existing travel agency websites, analyzing tourist needs, and studying modern web application architectures.', 'Requirement Collection:')
    add_bullet(doc, 'System Design: The system will be designed using DFDs (Data Flow Diagrams) for process mapping, ER (Entity-Relationship) diagrams for database structure, and UML Use Cases for defining actor (User, Admin, Chatbot) interactions.', 'System Design:')

    # 9. Implementation Tools
    add_heading(doc, '9. Implementation Tools')
    add_bullet(doc, 'Frontend: HTML5, CSS3, JavaScript, Bootstrap.', 'Frontend:')
    add_bullet(doc, 'Backend: Python (Flask).', 'Backend:')
    add_bullet(doc, 'Database: SQLite / MySQL.', 'Database:')
    add_bullet(doc, 'Machine Learning & AI: scikit-learn (K-Means, Linear Regression, TF-IDF Vectorizer), Google Gemini API.', 'Machine Learning & AI:')

    # 10. Expected Outcomes
    add_heading(doc, '10. Expected Outcomes')
    add_paragraph(doc, 'The tangible deliverables of the project include a fully functional web application deployed locally, a secure database for managing users and bookings, an interactive AI chatbot, and a robust admin dashboard featuring user-friendly reports and data visualizations.')

    # 11. Work Schedule
    add_heading(doc, '11. Work Schedule / Gantt Chart')
    add_bullet(doc, 'Week 1: Proposal Defense & Literature Review', 'Week 1:')
    add_bullet(doc, 'Week 2: System Design (ER, DFD, UML) & Prototyping', 'Week 2:')
    add_bullet(doc, 'Week 3-4: Frontend & Backend Coding (Flask, Database Integration)', 'Week 3-4:')
    add_bullet(doc, 'Week 5: Machine Learning Models & AI Chatbot Integration', 'Week 5:')
    add_bullet(doc, 'Week 6: System Testing & Debugging', 'Week 6:')
    add_bullet(doc, 'Week 7: Final Documentation & Project Presentation', 'Week 7:')

    doc.save(output_file)
    print(f"Docx file generated successfully at {output_file}")

if __name__ == "__main__":
    generate_proposal("Project_Proposal.docx")
